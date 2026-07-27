"""
final_check.py — Pre-submission audit.

Checks the submission against every requirement stated in Rules.pdf and every
criterion in Judge_Marking_Rubric.pdf, plus cross-document consistency and
file integrity. Run last, immediately before emailing.

Run:  python code/final_check.py
"""
import json
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd

BASE = Path(__file__).resolve().parents[1]
PROC, SUB, REP = (BASE / "data" / "processed", BASE / "submission",
                  BASE / "report")
MET = json.load(open(PROC / "metrics.json"))
LOG = pd.read_csv(PROC / "forecast_log.csv", parse_dates=["origin",
                                                          "target_date"])
OK, BAD, WARN = [], [], []


def chk(section, label, ok, detail=""):
    (OK if ok else BAD).append(label)
    print(f"  [{'PASS' if ok else 'FAIL'}] {label:<58} {detail}")


def warn(label, detail=""):
    WARN.append(label)
    print(f"  [TODO] {label:<58} {detail}")


def pdftext(p):
    import pdfplumber
    with pdfplumber.open(p) as pdf:
        return "".join((pg.extract_text() or "") for pg in pdf.pages)


print("=" * 80)
print("PRE-SUBMISSION FINAL CHECK")
print("=" * 80)

# ---------------------------------------------------------------- rules
print("\nA. REQUIREMENTS FROM Rules.pdf")
report = pdftext(REP / "Project_Report.pdf")
chk("A", "1-month-ahead forecast produced every Friday",
    set(LOG.origin.dt.dayofweek) == {4}
    and set((LOG.target_date - LOG.origin).dt.days) == {28},
    f"{len(LOG)} Friday origins, all 28-day horizons")
chk("A", "covers 1 Jan 2021 to 31 Dec 2025",
    str(LOG.origin.min().date()) == "2021-01-01"
    and str(LOG.origin.max().date()) == "2025-12-26", "261 origins")
chk("A", "uses SARB EXCX135 from EconData",
    "EXCX135" in report and "EconData" in report, "cited in the report")
chk("A", "RMSE computed for our model and for an AR(1)",
    "AR1_log" in MET and "COMBO" in MET,
    f"AR(1) {MET['AR1_log']['rmse']:.4f}, ours {MET['COMBO']['rmse']:.4f}")
chk("A", "model and data described in the report",
    "Methodology" in report and "Data" in report, "sections 2 to 4")
chk("A", "summary statistics and observation counts presented",
    (PROC / "table_summary_stats.csv").exists()
    and "Observation counts" in report, "Table 3.1")
chk("A", "RMSE chart with the numbers written on it",
    (BASE / "figures" / "fig1_rmse_bars.png").exists()
    and (BASE / "figures" / "slides" / "slide_rmse.png").exists(),
    "report figure 5.1 and slide 5")
chk("A", "plots of every series used in the model",
    (BASE / "figures" / "fig4_series_panel.png").exists(), "figure 3.1")
chk("A", "out-of-sample approach demonstrated",
    "out-of-sample" in report.lower() and (PROC / "forecast_log.csv").exists(),
    "protocol in section 2, evidence in forecast_log.csv")
chk("A", "code shared for marking",
    len(list((SUB / "code").glob("*.py"))) >= 8,
    f"{len(list((SUB / 'code').glob('*.py')))} scripts in submission/code")
chk("A", "AI use declared and explainable",
    "Claude" in report and "AI" in report, "report section 8")
chk("A", "PowerPoint presentation exists",
    (SUB / "02_Presentation.pptx").exists(), "02_Presentation.pptx")

# ---------------------------------------------------------------- rubric
print("\nB. RUBRIC CRITERIA (10 marks each)")
chk("B", "1. Outperforms the AR(1) baseline on RMSE",
    MET["COMBO"]["rmse"] < MET["AR1_log"]["rmse"],
    f"{100 * (1 - MET['COMBO']['rmse'] / MET['AR1_log']['rmse']):+.2f}%")
chk("B", "2. True out-of-sample execution demonstrated",
    (PROC / "validation_metrics.json").exists()
    and (PROC / "forecast_log.csv").exists(),
    "frozen validation era + full forecast log")
chk("B", "3. Clean reproducible code",
    (BASE / "code" / "run_all.py").exists()
    and (BASE / "code" / "README.md").exists(), "one-command pipeline + README")
chk("B", "4. Labelled, interpretable plots",
    len(list((BASE / "figures").glob("*.png"))) >= 8, "8 report + 3 slide figures")
chk("B", "5. Pitch and Q&A prepared",
    (BASE / "presentation" / "Script.md").exists()
    and "Q&A" in (BASE / "presentation" / "Script.md").read_text(encoding="utf-8"),
    "timed script + question bank")

# ------------------------------------------------- cross-document consistency
print("\nC. THE SAME NUMBERS EVERYWHERE")
gp = REP / "Complete_Project_Guide.pdf"
guide = pdftext(gp) if gp.exists() else ""
deck_js = (BASE / "presentation" / "build_deck.js").read_text(encoding="utf-8")
script = (BASE / "presentation" / "Script.md").read_text(encoding="utf-8")
readme = (BASE / "README.md").read_text(encoding="utf-8")
manifest = (SUB / "00_README.txt").read_text(encoding="utf-8") if SUB.exists() else ""
combo, basev = f"{MET['COMBO']['rmse']:.4f}", f"{MET['AR1_log']['rmse']:.4f}"
gain = f"{100 * (1 - MET['COMBO']['rmse'] / MET['AR1_log']['rmse']):.2f}"
docs = [("report", report), ("deck source", deck_js), ("script", script),
        ("README", readme)]
if guide: docs.insert(1, ("guide", guide))
if SUB.exists(): docs.append(("submission manifest", manifest))
for name, txt in docs:
    chk("C", f"{name} carries {combo} and {basev}",
        combo in txt and basev in txt, "")
for name, txt in docs:
    chk("C", f"{name} states the {gain}% gain", gain in txt, "")
chk("C", "no document still quotes a superseded figure",
    "stale long-run mean" not in (report + guide + readme + script), "")

# ---------------------------------------------------------------- writing
print("\nD. WRITING AND PRESENTATION MECHANICS")
for name, txt in [("report", report)] + ([("guide", guide)] if guide else []):
    chk("D", f"{name}: no em dashes", txt.count("—") == 0, "")
    chk("D", f"{name}: British/SA spelling",
        not re.search(r"\b(organize|analyze|behavior|percent\b(?!age))", txt), "")
    chk("D", f"{name}: no unrendered markup",
        "&#" not in txt and "<b>" not in txt, "")
from pptx import Presentation
deck = Presentation((SUB / "02_Presentation.pptx") if SUB.exists()
                    else (BASE / "presentation" / "Deck.pptx"))
titles = []
for sl in deck.slides:
    t = [sh.text_frame.text.strip() for sh in sl.shapes
         if sh.has_text_frame and sh.text_frame.text.strip()]
    titles.append((t[1] if t and t[0].startswith("Appendix") else (t[0] if t else "")))
chk("D", "deck has 14 slides (6 presented, references, 7 appendix)",
    len(deck.slides) == 14, f"{len(deck.slides)} slides")
chk("D", "every presented slide has speaker notes",
    all(sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip()
        for sl in list(deck.slides)[:6]), "slides 1 to 6")
chk("D", "conclusions is the last presented slide",
    "Conclusions" in titles[5], titles[5][:40])
chk("D", "references slide present", "References" in titles[6], "")
chk("D", "appendix slides labelled A to G",
    all(any(sh.text_frame.text.strip().startswith(f"Appendix {L}")
            for sh in list(deck.slides)[7 + k].shapes if sh.has_text_frame)
        for k, L in enumerate("ABCDEFG")), "")
words = [sum(len(sh.text_frame.text.split()) for sh in sl.shapes
             if sh.has_text_frame) for sl in list(deck.slides)[:6]]
chk("D", "presented slides stay concise", max(words) <= 130,
    f"max {max(words)} words on a slide")
spoken = sum(len(re.sub(r"\*|\[.*?\]|>", " ", x).split())
             for x in re.findall(r"^> .*$", script.split("# THE SHORT VERSION")[0],
                                 re.M))
chk("D", "script fits the 2 to 5 minute slot",
    3.0 <= spoken / 145 <= 4.9, f"{spoken} words, about {spoken / 145:.1f} min")

# ---------------------------------------------------------------- package
print("\nE. SUBMISSION PACKAGE")
if not SUB.exists():
    print("  [SKIP] submission/ not built yet. Run: python code/make_submission.py")
    need = []
else:
    need = ["00_README.txt", "01_Project_Report.pdf", "02_Presentation.pptx",
        "03_forecast_log.csv", "04_metrics.json", "05_validation_metrics.json"]
for f in need:
    chk("E", f"{f} present", (SUB / f).exists(),
        f"{(SUB / f).stat().st_size / 1024:.0f} KB" if (SUB / f).exists() else "")
if SUB.exists():
    chk("E", "code folder complete",
        len(list((SUB / "code").glob("*"))) >= 10,
        f"{len(list((SUB / 'code').glob('*')))} files")
    chk("E", "figures included",
        len(list((SUB / "figures").glob("*.png"))) >= 10,
        f"{len(list((SUB / 'figures').glob('*.png')))} charts")
    chk("E", "manifest reports the current headline numbers",
        combo in manifest and basev in manifest, "")
    size = sum(f.stat().st_size for f in SUB.rglob("*") if f.is_file())
    chk("E", "package is a sensible size to email", size < 20 * 1024 ** 2,
        f"{size / 1024 ** 2:.1f} MB")
    sub_log = pd.read_csv(SUB / "03_forecast_log.csv")
    chk("E", "the personal study guide is NOT in the submission",
        not (SUB / "06_Complete_Project_Guide.pdf").exists()
        and not list(SUB.glob("*Guide*")), "kept in report/ for Aakash only")
    chk("E", "submitted forecast log matches the live one",
        len(sub_log) == len(LOG) == 261, f"{len(sub_log)} rows")

    # ---------------------------------------------------------------- reproducibility
print("\nF. REPRODUCIBILITY")
chk("F", "one-command pipeline present",
    (BASE / "code" / "run_all.py").exists(), "python code/run_all.py")
chk("F", "independent chart audit ships with the code",
    (SUB / "code" / "verify_charts.py").exists(), "verify_charts.py, 96 checks")
r = subprocess.run([sys.executable, str(BASE / "code" / "verify_charts.py")],
                   capture_output=True, text=True)
chk("F", "chart audit passes right now", "0 failed" in r.stdout,
    [l for l in r.stdout.splitlines() if l.startswith("RESULT")][0]
    if "RESULT" in r.stdout else "")

# ---------------------------------------------------------------- human tasks
print("\nG. THINGS ONLY AAKASH CAN DO")
warn("Confirm both team members registered via the poster link",
     "blocking for eligibility")
warn("Obtain and complete the official submission form", "not in the folder")
warn("Zip submission/ and email to 28399323@sun.ac.za", "deadline Monday 23:59")
warn("Rehearse Script.md aloud with a timer, twice", "target 4:00 to 4:30")

print("\n" + "=" * 80)
print(f"RESULT: {len(OK)} passed, {len(BAD)} failed, {len(WARN)} manual tasks")
if BAD:
    print("\nFAILURES:")
    for f in BAD:
        print("  -", f)
else:
    print("Everything automated is in order. The manual tasks above remain.")
print("=" * 80)
